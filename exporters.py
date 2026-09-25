"""Create plain-text, Word, and PDF versions of an editable draft."""
from io import BytesIO
from pathlib import Path
import re

LOGO_PATH = Path(__file__).resolve().parent.parent / "assets" / "legalease_logo.png"


def make_txt(content: str, document_type: str) -> bytes:
    return f"{document_type.upper()}\n\n{content.strip()}\n".encode("utf-8")


def _is_heading(line: str) -> bool:
    clean = line.strip()
    return bool(clean) and (clean.isupper() or bool(re.match(r"^\d+(?:\.\d+)*[.)]?\s+", clean)))


def make_docx(content: str, document_type: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.add_run().add_picture(str(LOGO_PATH), width=Inches(2.0))
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(document_type.upper())
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(28, 51, 71)
    for line in content.splitlines():
        if not line.strip():
            continue
        p = doc.add_paragraph()
        if _is_heading(line):
            run = p.add_run(line.strip())
            run.bold = True
            run.font.color.rgb = RGBColor(28, 51, 71)
        else:
            p.add_run(line.strip())
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase  |  Draft for review - not legal advice").italic = True
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def make_pdf(content: str, document_type: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    from xml.sax.saxutils import escape

    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=LETTER, rightMargin=0.85 * inch,
                            leftMargin=0.85 * inch, topMargin=0.85 * inch, bottomMargin=0.75 * inch)
    base = getSampleStyleSheet()
    title = ParagraphStyle("LegalEaseTitle", parent=base["Title"], fontName="Helvetica-Bold",
                           fontSize=16, leading=20, textColor=colors.HexColor("#1c3347"), alignment=TA_CENTER, spaceAfter=20)
    body = ParagraphStyle("LegalEaseBody", parent=base["BodyText"], fontName="Times-Roman", fontSize=10.5,
                          leading=15, spaceAfter=8, textColor=colors.HexColor("#253746"))
    heading = ParagraphStyle("LegalEaseHeading", parent=body, fontName="Helvetica-Bold", fontSize=10.5,
                             textColor=colors.HexColor("#1c3347"), spaceBefore=7, keepWithNext=True)
    flow = [Paragraph(escape(document_type.upper()), title)]
    for line in content.splitlines():
        if line.strip():
            flow.append(Paragraph(escape(line.strip()), heading if _is_heading(line) else body))
        else:
            flow.append(Spacer(1, 4))

    def footer(canvas, _document):
        canvas.saveState()
        canvas.drawImage(ImageReader(str(LOGO_PATH)), 0.85 * inch, LETTER[1] - 0.62 * inch,
                         width=1.35 * inch, height=0.28 * inch, preserveAspectRatio=True, mask="auto")
        canvas.setStrokeColor(colors.HexColor("#d8e0e5"))
        canvas.line(0.85 * inch, 0.58 * inch, LETTER[0] - 0.85 * inch, 0.58 * inch)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#61717d"))
        canvas.drawString(0.85 * inch, 0.4 * inch, "LegalEase | Draft for review - not legal advice")
        canvas.drawRightString(LETTER[0] - 0.85 * inch, 0.4 * inch, f"Page {doc.page}")
        canvas.restoreState()

    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
